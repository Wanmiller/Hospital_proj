"""Build offline Kazakh defense slides and handouts. No database writes.

Build tools only: Python, python-docx, markdown, Playwright, pypdf.
Run from any directory: python docs/defense/build.py
"""
from pathlib import Path
from html import escape
import json
import re
import shutil
import zipfile
import markdown
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from playwright.sync_api import sync_playwright
from pypdf import PdfReader

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
IMAGES = OUT/'images'
IMAGES.mkdir(exist_ok=True)
for name, source in {
    'patient-booking.png': ROOT/'docs/patient-preview/booking.png',
    'consultation.png': ROOT/'docs/tabler-preview/consultation.png',
}.items():
    shutil.copyfile(source, IMAGES/name)

def cards(items):
    return '<div class="cards">'+''.join(f'<article><h3>{escape(h)}</h3><p>{escape(t)}</p></article>' for h,t in items)+'</div>'

def bullets(items):
    return '<ul>'+''.join(f'<li>{escape(item)}</li>' for item in items)+'</ul>'

slides = [
    ('1-адам · 00:00–00:40', 'Емханадағы қабылдауға жазылу жүйесі',
     '<p class="lead">Пациенттің жазылуынан қабылдаудың нәтижесіне дейін</p>'
     +cards([('4 рөл','Пациент · тіркеуші · дәрігер · әкімші'),('9 негізгі кесте','Өзара байланысқан деректер'),('15 минут','12 минут баяндама және көрсетілім + 3 минут сұрақ')])
     +'<p class="muted">Оқу жобасы · PHP + MySQL/MariaDB · Қазақша интерфейс</p>',
     'Мақсатты айтыңыз: қабылдауға жазу, медициналық нәтижені тіркеу, есеп алу. Аты-жөндеріңізді ауызша таныстырыңыздар.'),
    ('1-адам · 00:40–01:30', 'Кім не істейді?',
     cards([('Пациент','Өзі тіркеледі, бос уақытқа жазылады, өз қабылдауларын көреді.'),('Тіркеуші','Пациент карточкаларын жүргізеді, қабылдауға жазады және жазбаны басқарады.'),('Дәрігер','Өз қабылдаулары, медициналық жазба, рецепт және пациент тарихы.'),('Әкімші','Аккаунттар, дәрігерлер, бөлімшелер, есептер және бар карточкаға аккаунт байланыстыру.')]),
     'Пациенттің өзі жазыла алатынын ерекше айтыңыз. Карточкасы бар әр пациентке аккаунт міндетті емес.'),
    ('1-адам · 01:30–02:20', 'Сұрау қалай орындалады?',
     '<div class="flow"><article>Браузер<small>HTML · Tabler · JavaScript</small></article><b>→</b><article>PHP сервері<small>Сессия · рөл · тексеріс</small></article><b>→</b><article>MySQL / MariaDB<small>PDO · кілттер · транзакция</small></article></div>'
     +bullets(['Нәтиже қазақша HTML беті ретінде қайтады.','Tabler Core 1.7.1 файлдары жергілікті сақталған.','Жергілікті іске қосу үшін PHP сервері мен база қажет.']),
     'Браузер базаға тікелей қосылмайды. PHP сұрауды тексеріп, параметрлері бар PDO сұрауын орындайды. GitHub кодты сақтайды, сайтты өзі іске қоспайды.'),
    ('1-адам · 02:20–04:10', 'Дерекқордың негізгі байланыстары',
     '<div class="schema"><div class="schema-row"><span>patients<br><small>Пациент</small></span><b>1 → N</b><span>appointments<br><small>Қабылдау</small></span><b>N ← 1</b><span>doctors<br><small>Дәрігер</small></span></div>'
     '<div class="schema-row"><span>appointments</span><b>1 → 0..1</b><span>medical_records</span><b>1 → N</b><span>prescriptions</span></div></div>'
     +cards([('users + patient_accounts','Аккаунт пен пациент карточкасының жеке байланысы.'),('departments + specializations','Бөлімше мен мамандық анықтамалықтары.')])
     +'<p class="muted">Медициналық жазбаға бірнеше рецепт берілуі мүмкін. Пациент карточкасы аккаунтсыз да болады.</p>',
     'Тоғыз негізгі кестені атаңыз. Бір қабылдауда бір пациент пен бір дәрігер бар. medical_records.appointment_id UNIQUE. doctors.user_id міндетті және бірегей. Қосымша зертханалық кестелер негізгі схемаға кірмейді.'),
    ('1-адам · 04:10–06:00', 'Тұтастықты тексеру',
     '<div class="two"><div>'+bullets(['PK — әр жазбаның идентификаторы.','FK — байланысты жазбаның бар екенін және рөлін тексереді.','UNIQUE — дәрігердің бір уақытын екі рет алуға жол бермейді.','Транзакция — медициналық жазба мен мәртебені бірге сақтайды.'])+'</div><div class="code-card"><h3>unique_doctor_slot</h3><pre>doctor_id\nappointment_date\nappointment_time\noccupied_slot</pre><p>Scheduled / Completed → 1<br>Cancelled → NULL</p></div></div>'
     +'<p class="muted">FOR UPDATE қабылдауды қатар аяқтау мен болдырмауды реттейді.</p>',
     'Cancelled слотты босатады, Completed бұрынғы орнын сақтайды. Индекс бірдей басталу уақытын қорғайды; ерікті ұзақтықтағы интервалдарды салыстыру жоқ. 06:00-де екінші баяндамашыға сөз беріңіз.'),
    ('2-адам · 06:00–06:40', 'Пациенттің өзі жазылуы',
     '<div class="two screenshot-layout"><div>'+bullets(['Тіркелу → аккаунт және пациент карточкасы.','Дәрігер мен күн → бос уақыттар.','Бүгіннен бастап 90 күнге дейін.','09:00–17:00 · әр 30 минут сайын.','Тек өз қабылдауларын көру және жоспарланған жазбадан бас тарту.'])+'</div><figure><img src="images/patient-booking.png" alt="Пациенттің дәрігер мен бос уақытты таңдау беті"><figcaption>Бөлек тест базасындағы жасанды деректер</figcaption></figure></div>',
     'Соңғы уақыт 16:30. Жеке дәрігер кестесі жоқ, демалыс күндері бөлек жүргізілмейді. Бұрынғы карточкаға аккаунтты әкімші байланыстырады.'),
    ('2-адам · 06:40–10:00', 'Көрсетілім: бір қабылдаудың жолы',
     '<div class="two screenshot-layout"><div>'+bullets(['Пациент: бос уақытқа жазылу.','Тіркеуші: сол күндегі жаңа жазбаны көру.','Дәрігер: медициналық жазбаны сақтау.','Пациент: «Аяқталған» мәртебесі.','Әкімші: күн аралығын таңдап есеп алу.'])+'</div><figure><img src="images/consultation.png" alt="Дәрігердің консультация беті"><figcaption>Алдын ала түсірілген оқу скриншоты</figcaption></figure></div>',
     'Сайтқа ауысыңыз. Төрт бөлек браузер профилі дайын болсын. Бір күн мен бір дәрігерді таңдаңыздар. Қате болса, уақыт жоғалтпай осы скриншоттарды бұрынғы оқу нәтижесі деп түсіндіріңіздер. 10:00-де слайдқа оралыңыз.'),
    ('2-адам · 10:00–10:50', 'Қолжетімділік серверде тексеріледі',
     cards([('Рөл және сессия','Әр қорғалған сұрауда аккаунттың белсенділігі мен рөлі тексеріледі.'),('Деректің иесі','Пациент карточкасы сессиядан анықталады; бөтен қабылдауды болдырмайды.'),('Формалар мен мәтін','POST формаларында CSRF; HTML шығару кезінде экрандау; SQL мәндері PDO параметрлерімен.'),('Құпиясөз','password_hash арқылы сақталады, password_verify арқылы тексеріледі.')]),
     'Мәзірді жасыру жеткіліксіз. Дәрігердің қабылдауға иелігі де серверде тексеріледі. Бұғатталған аккаунт келесі қорғалған сұрауда шығарылады.'),
    ('2-адам · 10:50–11:30', 'Тексеру нәтижелері және шектеулер',
     '<div class="two"><div><h3>Сақталған браузерлік нәтижелер</h3><div class="metrics"><article><strong>57</strong><span>бұрынғы рөлдер</span></article><article><strong>38</strong><span>пациент сценарийлері</span></article></div><p class="muted">27.10.2026 нәтижелері · бөлек тест базасы</p><p>Рөлдер, қақтығыс, иелік, мобильді мәзір және пернетақта.</p></div><div><h3>Келесі дамыту қадамдары</h3>'+bullets(['Жеке дәрігер кестесі, демалыс және үзіліс.','SMS немесе басқа еске салу қызметі.','Пациентке медициналық жазбаларды көрсету.','Өзгерістер аудиті және орналастыру қорғаныстары.'])+'</div></div>',
     'Бұл сандар JSON нәтижелерінен алынған; бүгін қайта орындалған тест ретінде көрсетпеңіз. Қазіргі нұсқада жоқ мүмкіндіктерді дайын функция деп атамаңыз.'),
    ('2-адам · 11:30–12:00 · кейін 3 минут сұрақ', 'Нәтиже',
     '<p class="lead">Жазылу → қабылдау → медициналық нәтиже → есеп</p>'
     +cards([('Пациент үшін','Өздігінен жазылу және жеке қабылдаулар тізімі.'),('Емхана үшін','Тіркеуші мен дәрігердің байланысқан жұмыс барысы.'),('Дерекқор үшін','Кілттер, бірегей уақыт және транзакциялар.')])
     +'<p class="repo">github.com/Wanmiller/Hospital_proj</p><p class="muted">Сұрақтарға дайынбыз. SQL мен кодқа арналған мысалдар guide.docx ішінде.</p>',
     'Екеуіңіз де сұрақтарға қатысыңыздар: бірінші адам — база, байланыс, индекс, транзакция; екінші адам — интерфейс, авторизация, пациент және демонстрация.'),
]

style = """
:root{--blue:#245ba8;--ink:#182b43;--muted:#586b81;--line:#dce4ee;--scale:1}
*{box-sizing:border-box}body{margin:0;background:#152237;font-family:'Segoe UI',Arial,sans-serif;color:var(--ink)}
main{position:absolute;left:50%;top:50%;width:1280px;height:720px;transform:translate(-50%,-50%) scale(var(--scale));transform-origin:center}
.slide{display:none;width:1280px;height:720px;padding:45px 62px 56px;background:#f7f9fc;position:relative;overflow:hidden;border-top:8px solid var(--blue)}
.slide.active{display:block}.eyebrow{font-size:18px;font-weight:600;color:var(--blue);margin-bottom:19px}h1{font-size:40px;line-height:1.14;margin:0 0 27px;letter-spacing:-.7px;max-width:1120px}h3{font-size:23px;margin:0 0 12px;line-height:1.25}p{font-size:23px;line-height:1.45;margin:12px 0 17px}.lead{font-size:31px;color:var(--blue);margin-bottom:35px}.muted{font-size:18px;color:var(--muted)}
.cards{display:flex;flex-wrap:wrap;gap:18px;margin:24px 0}.cards article{flex:1 1 30%;padding:23px;background:#fff;border:1px solid var(--line);border-radius:8px}.cards:has(article:nth-child(4)) article{flex-basis:46%}.cards p{font-size:21px;margin:0}.two{display:grid;grid-template-columns:1fr 1fr;gap:36px}.two ul{margin-top:5px}ul{padding-left:26px;margin:24px 0}li{font-size:24px;line-height:1.4;margin-bottom:18px}.code-card{background:#eaf0f8;padding:25px 32px;border-radius:8px}.code-card pre{font-size:25px;line-height:1.4;margin:16px 0}.code-card p{font-size:20px}.flow{display:flex;gap:20px;align-items:center;margin:48px 0 28px}.flow article{background:white;border:1px solid var(--line);padding:28px 20px;flex:1;text-align:center;font-size:27px;font-weight:600;border-radius:7px}.flow b{color:var(--blue);font-size:36px}.flow small{display:block;font-size:18px;font-weight:400;margin-top:16px}
.schema{padding:5px 0}.schema-row{display:flex;align-items:center;justify-content:space-between;gap:14px;margin:17px 0}.schema-row span{font-size:23px;text-align:center;flex:1;padding:17px 5px;border:1px solid var(--line);background:white;border-radius:7px}.schema-row small{font-size:18px;color:var(--muted)}.schema-row b{font-size:20px;color:var(--blue)}
figure{margin:0;background:#fff;border:1px solid var(--line);padding:10px;border-radius:7px;height:440px;display:flex;flex-direction:column}figure img{width:100%;min-height:0;flex:1;object-fit:contain;object-position:top}figcaption{font-size:14px;color:var(--muted);text-align:center;margin-top:8px}.screenshot-layout{grid-template-columns:.8fr 1.2fr}.screenshot-layout li{font-size:23px}.metrics{display:flex;gap:16px}.metrics article{background:#eaf0f8;border-radius:6px;padding:22px;flex:1}.metrics strong{display:block;font-size:60px;color:var(--blue)}.metrics span{font-size:19px}.repo{font-size:26px;color:var(--blue);margin-top:30px}
footer{position:absolute;bottom:20px;left:62px;right:62px;display:flex;justify-content:space-between;font-size:15px;color:var(--muted);border-top:1px solid var(--line);padding-top:12px}.notes{display:none;position:absolute;bottom:48px;left:30px;right:30px;background:#14233e;color:#fff;padding:18px;font-size:21px;line-height:1.4;z-index:5;border-radius:6px}.show-notes .active .notes{display:block}.controls{position:fixed;bottom:6px;left:0;right:0;display:flex;justify-content:center;gap:12px;color:white;font:14px 'Segoe UI',sans-serif}.controls button{border:1px solid #7d94b3;border-radius:4px;background:#263f60;color:white;padding:5px 14px;cursor:pointer}.controls span{align-self:center}
@media print{@page{size:1280px 720px;margin:0}body{background:white}main{position:static;transform:none;width:1280px;height:auto}.slide{display:block!important;break-after:page;page-break-after:always;print-color-adjust:exact;-webkit-print-color-adjust:exact}.slide:last-child{break-after:auto;page-break-after:auto}.controls,.notes{display:none!important}}
"""
sections = []
for i,(speaker,title,body,note) in enumerate(slides,1):
    sections.append(f'<section class="slide{" active" if i==1 else ""}" aria-label="{i}-слайд"><div class="eyebrow">{escape(speaker)}</div><h1>{escape(title)}</h1>{body}<aside class="notes">{escape(note)}</aside><footer><span>Емхана · Екі адамға арналған қорғау</span><span>{i} / {len(slides)}</span></footer></section>')
script = """
let index=0;const pages=[...document.querySelectorAll('.slide')];
function go(delta){index=Math.max(0,Math.min(pages.length-1,index+delta));pages.forEach((page,i)=>page.classList.toggle('active',i===index));document.getElementById('counter').textContent=`${index+1} / ${pages.length}`;}
function resize(){document.documentElement.style.setProperty('--scale',Math.min(innerWidth/1280,(innerHeight-35)/720));}
addEventListener('resize',resize);resize();
document.addEventListener('keydown',event=>{if(['ArrowRight','PageDown',' '].includes(event.key)){event.preventDefault();go(1);}if(['ArrowLeft','PageUp'].includes(event.key)){event.preventDefault();go(-1);}if(event.key.toLowerCase()==='n')document.body.classList.toggle('show-notes');if(event.key.toLowerCase()==='f'){if(document.fullscreenElement)document.exitFullscreen();else document.documentElement.requestFullscreen().catch(()=>{});}if(event.key==='Home'){index=0;go(0);}if(event.key==='End'){index=pages.length-1;go(0);}});
document.getElementById('previous').onclick=()=>go(-1);document.getElementById('next').onclick=()=>go(1);
"""
slides_html='<!doctype html><html lang="kk"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Емхана — 15 минуттық қорғау</title><style>'+style+'</style></head><body><main>'+''.join(sections)+'</main><nav class="controls" aria-label="Слайдтарды ауыстыру"><button id="previous">← Алдыңғы</button><span id="counter">1 / 10</span><button id="next">Келесі →</button><span>F: толық экран · N: ескертпе</span></nav><script>'+script+'</script></body></html>'
(OUT/'slides.html').write_text(slides_html,encoding='utf-8')

guide = (OUT/'guide-kk.md').read_text(encoding='utf-8')
guide_body = markdown.markdown(guide,extensions=['tables','fenced_code'])
guide_style = """
@page{size:A4;margin:17mm 17mm 18mm}*{box-sizing:border-box}body{font:11pt/1.5 'Segoe UI',Arial,sans-serif;color:#182b43}h1{font-size:22pt;color:#245ba8}h2{font-size:16pt;color:#245ba8;margin-top:24px}h3{font-size:12pt;margin-top:19px}h1,h2,h3{break-after:avoid}p{orphans:3;widows:3}table{border-collapse:collapse;width:100%;font-size:9pt;margin:12px 0}th,td{border:1px solid #ccd8e6;padding:7px;text-align:left;vertical-align:top}th{background:#eaf0f8}thead{display:table-header-group}tr{break-inside:avoid}pre{background:#eef2f7;border:1px solid #dae2ed;padding:10px;font:9pt/1.45 Consolas,monospace;white-space:pre-wrap;overflow-wrap:anywhere;break-inside:avoid}code{font:9.5pt Consolas,monospace}a{color:#245ba8}li{margin-bottom:6px}
"""
(OUT/'guide.html').write_text('<!doctype html><html lang="kk"><head><meta charset="utf-8"><title>Екі адамға арналған қорғау материалы</title><style>'+guide_style+'</style></head><body>'+guide_body+'</body></html>',encoding='utf-8')
cheatsheet=markdown.markdown((OUT/'cheatsheet.md').read_text(encoding='utf-8'))
cheat_style=guide_style+'@page{margin:12mm}body{font-size:10.5pt;line-height:1.3}h1{font-size:19pt;margin:0 0 8px}h2{font-size:12pt;margin:12px 0 5px}p{margin:5px 0}ul{padding-left:19px;margin:4px 0}li{margin-bottom:4px}code{font-size:9pt}'
(OUT/'cheatsheet.html').write_text('<!doctype html><html lang="kk"><head><meta charset="utf-8"><title>Қорғау жадынамасы</title><style>'+cheat_style+'</style></head><body>'+cheatsheet+'</body></html>',encoding='utf-8')

# Word version uses the same Markdown source, including tables and code blocks.
doc=Document()
section=doc.sections[0]
section.top_margin=section.bottom_margin=Cm(1.8)
section.left_margin=section.right_margin=Cm(1.8)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6)
for key in ['Title','Heading 1','Heading 2','Heading 3']:
    doc.styles[key].font.name='Calibri';doc.styles[key].font.color.rgb=RGBColor.from_string('245BA8')

def clean(text):
    return re.sub(r'\[([^]]+)\]\([^)]+\)',r'\1',text).replace('**','').replace('`','')

lines=guide.splitlines();i=0;in_code=False
while i<len(lines):
    line=lines[i]
    if line.startswith('```'):
        in_code=not in_code;i+=1;continue
    if in_code:
        paragraph=doc.add_paragraph();paragraph.paragraph_format.space_after=Pt(0)
        run=paragraph.add_run(line);run.font.name='Consolas';run.font.size=Pt(8)
    elif line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            cells=[clean(c.strip()) for c in lines[i].strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',c) for c in cells): rows.append(cells)
            i+=1
        table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Light Shading Accent 1'
        for n,row in enumerate(rows):
            cells=table.add_row().cells
            for c,value in zip(cells,row):
                c.text=value
                if n==0:
                    for run in c.paragraphs[0].runs: run.bold=True
        continue
    elif line.startswith('# '): doc.add_heading(clean(line[2:]),0)
    elif line.startswith('## '): doc.add_heading(clean(line[3:]),1)
    elif line.startswith('### '): doc.add_heading(clean(line[4:]),2)
    elif line.strip():
        paragraph=doc.add_paragraph(clean(line))
        if line.startswith('**'): paragraph.runs[0].bold=True
    i+=1
section.footer.paragraphs[0].text='Емханадағы қабылдауға жазылу жүйесі · 15 минут · 2 адам'
doc.save(OUT/'guide.docx')

errors=[]; external=[]
with sync_playwright() as p:
    browser=p.chromium.launch()
    page=browser.new_page(viewport={'width':1280,'height':760},device_scale_factor=1)
    page.on('pageerror',lambda e: errors.append(str(e)))
    page.on('request',lambda req: external.append(req.url) if req.url.startswith(('https://','http://')) else None)
    page.goto((OUT/'slides.html').as_uri(),wait_until='networkidle')
    page.evaluate('document.fonts.ready')
    assert page.locator('.slide').count()==10
    assert page.locator('img').evaluate_all('(images)=>images.every(img=>img.complete && img.naturalWidth>0)')
    # Verify every text/image box fits inside its slide, including hidden slides.
    page.add_style_tag(content='.slide{display:block!important}main{position:static;transform:none;height:auto}')
    bounds=page.locator('.slide').evaluate_all('''slides=>slides.map((slide,index)=>{
      const rect=slide.getBoundingClientRect();
      const bad=[...slide.querySelectorAll('h1,h3,p,li,article,figure,.schema-row')].filter(el=>{const r=el.getBoundingClientRect();return r.bottom>rect.bottom-45||r.right>rect.right-25||r.left<rect.left;});
      return {slide:index+1,overflow:bad.map(el=>el.textContent.slice(0,70))};
    })''')
    assert not [row for row in bounds if row['overflow']],bounds
    page.pdf(path=str(OUT/'slides.pdf'),print_background=True,prefer_css_page_size=True)
    # Capture a contact sheet input for visual review, one real screenshot per slide.
    preview=ROOT/'tests/.runtime/defense-preview';preview.mkdir(parents=True,exist_ok=True)
    for idx,slide in enumerate(page.locator('.slide').all(),1):
        slide.screenshot(path=str(preview/f'{idx:02}.png'))
    page.goto((OUT/'slides.html').as_uri(),wait_until='networkidle')
    page.keyboard.press('ArrowRight')
    assert page.locator('.slide.active .eyebrow').inner_text().startswith('1-адам')
    assert page.locator('#counter').inner_text()=='2 / 10'
    page.keyboard.press('End');assert page.locator('#counter').inner_text()=='10 / 10'
    page.keyboard.press('n');assert page.locator('.active .notes').is_visible()
    page.keyboard.press('Home');assert page.locator('#counter').inner_text()=='1 / 10'
    page.goto((OUT/'guide.html').as_uri(),wait_until='networkidle')
    page.pdf(path=str(OUT/'guide.pdf'),print_background=True,prefer_css_page_size=True,display_header_footer=True,header_template='<span></span>',footer_template='<div style="font-family:Arial;font-size:8px;width:100%;text-align:center;color:#65758b"><span class="pageNumber"></span> / <span class="totalPages"></span></div>')
    page.goto((OUT/'cheatsheet.html').as_uri(),wait_until='networkidle')
    page.pdf(path=str(OUT/'cheatsheet.pdf'),print_background=True,prefer_css_page_size=True)
    browser.close()

assert not errors,errors
assert not external,external
slide_pdf=PdfReader(OUT/'slides.pdf')
guide_pdf=PdfReader(OUT/'guide.pdf')
assert len(slide_pdf.pages)==10
assert len(PdfReader(OUT/'cheatsheet.pdf').pages)==1
for number,page in enumerate(slide_pdf.pages,1):
    assert len(page.extract_text().strip())>90,number
guide_text='\n'.join(page.extract_text() for page in guide_pdf.pages)
for term in ['1-адам','2-адам','12:00–15:00','patient_accounts','occupied_slot','Ықтимал сұрақтар']:
    assert term in guide_text,term
assert len(Document(OUT/'guide.docx').tables)==5
result={'date':'2026-10-28','slides':10,'guide_pdf_pages':len(guide_pdf.pages),'cheatsheet_pdf_pages':1,'word_tables':5,'presentation_minutes':12,'questions_minutes':3,'total_minutes':15,'javascript_errors':errors,'external_requests':external,'slide_bounds':bounds,'source':'Current PHP/SQL and stored browser verification results dated 2026-10-27'}
(OUT/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
with zipfile.ZipFile(OUT/'defense-materials.zip','w',zipfile.ZIP_DEFLATED) as bundle:
    for name in ['README.md','guide-kk.md','guide.docx','guide.html','guide.pdf','slides.html','slides.pdf','cheatsheet.md','cheatsheet.html','cheatsheet.pdf','verification.json']:
        bundle.write(OUT/name,name)
    for file in sorted(IMAGES.glob('*.png')): bundle.write(file,'images/'+file.name)
print(json.dumps(result,ensure_ascii=False,indent=2))
